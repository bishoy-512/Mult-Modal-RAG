from controllers import BaseController
from .Enums import ExtensionsEnum
from langchain_community.document_loaders import TextLoader , PyMuPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from paddleocr import PaddleOCR
from docx import Document as DocxDocument
from pptx import Presentation
import time


class ProcessController(BaseController):
    def __init__(self):
        super().__init__()

    ocr = None

    def get_ocr(self):
        if ProcessController.ocr is None:
            ProcessController.ocr = PaddleOCR(
                lang="ar",
                enable_mkldnn=False,
                use_doc_orientation_classify=False,
                use_doc_unwarping=False,
                use_textline_orientation=True,
            )
        return ProcessController.ocr
        
    def process_file(self, file_path : str, chunk_size : int = 500, overlap_size : int = 50):
        file_ext = "." + file_path.split(".")[-1]
                
        if file_ext == ExtensionsEnum.TXT.value:
            return self.process_txt_file(file_path=file_path,chunk_size=chunk_size,overlap_size=overlap_size)
        
        if file_ext == ExtensionsEnum.PDF.value:
            return self.process_pdf_file(file_path=file_path,chunk_size=chunk_size,overlap_size=overlap_size)
        
        if file_ext in ExtensionsEnum.IMAGE.value:
            return self.process_image_file(file_path=file_path,chunk_size=chunk_size,overlap_size=overlap_size)
        
        if file_ext == ExtensionsEnum.WORD.value:
            return self.process_word_file(file_path=file_path,chunk_size=chunk_size,overlap_size=overlap_size)
            
        if file_ext == ExtensionsEnum.POWERPOINT.value:
            return self.process_pptx_file(file_path=file_path,chunk_size=chunk_size,overlap_size=overlap_size)
            
        return None
    
    def process_txt_file(self,file_path: str,chunk_size: int = 500,overlap_size: int = 50):
        loader = TextLoader(file_path=file_path,encoding="utf-8")
        content = loader.load()
        return self.process_file_content(content=content,chunk_size=chunk_size,overlap_size=overlap_size)
    
    def process_pdf_file(self,file_path: str,chunk_size: int = 500,overlap_size: int = 50):
        loader = PyMuPDFLoader(file_path=file_path)
        content = loader.load()
        return self.process_file_content(content=content,chunk_size=chunk_size,overlap_size=overlap_size)
    
    def process_image_file(self, file_path: str,chunk_size: int = 500,overlap_size: int = 50):
        
        start = time.time()
        ocr = self.get_ocr()
        end = time.time()
        print(f"OCR init time: {end - start:.2f} seconds")
        
        start = time.time()
        result = ocr.predict(file_path)
        end = time.time()
        print(f"OCR predict time: {end - start:.2f} seconds")
        
        extracted_text = []
        for res in result:
            extracted_text.extend(res["rec_texts"])

        text = "\n".join(extracted_text)
        content = [
            Document(
                page_content=text,
                metadata={
                    "source": file_path,
                    "type": "image",
                    "ocr": True
                }
            )
        ]
        return self.process_file_content(content=content,chunk_size=chunk_size,overlap_size=overlap_size)

    def process_word_file(self, file_path: str,chunk_size: int = 500,overlap_size: int = 50):
        doc = DocxDocument(file_path)
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        content = [
            Document(
                page_content="\n".join(text),
                metadata={
                    "source": file_path,
                    "type": "docx"
                }
            )
        ]
        return self.process_file_content(content=content,chunk_size=chunk_size,overlap_size=overlap_size)

    def process_pptx_file(self,file_path: str,chunk_size: int = 500,overlap_size: int = 50):
        presentation = Presentation(file_path)
        
        text = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                if shape.text.strip():
                    text.append(shape.text)
                    
        text = "\n".join(text)
        content = [
            Document(
                page_content=text,
                metadata={
                    "source": file_path,
                    "type": "pptx"
                }
            )
        ]
        return self.process_file_content(content=content,chunk_size=chunk_size,overlap_size=overlap_size)


    def process_file_content(self , content : list , chunk_size : int = 500 , overlap_size : int = 50):
        text_splitter_model = RecursiveCharacterTextSplitter(chunk_size = chunk_size, chunk_overlap = overlap_size)
        return text_splitter_model.split_documents(content)
